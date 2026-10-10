"""Resume import: extract text locally, then (optionally) structure it with the user's LLM."""

from __future__ import annotations

import io
import json

from pydantic import ValidationError

from kestrel.llm.provider import ChatClient, LLMConfig
from kestrel.profile.schema import CandidateProfile


def extract_text(filename: str, data: bytes) -> str:
    """Plain text of a resume. Raises ValueError with a readable message for unusable files."""
    name = filename.lower()
    if not name.endswith((".pdf", ".docx", ".txt", ".md")):
        raise ValueError("Supported resume formats: .pdf, .docx, .txt, .md")
    try:
        text = _extract(name, data)
    except Exception as e:  # corrupt or unusual files: explain instead of failing
        raise ValueError(
            f"Couldn't read this file ({type(e).__name__}). Try saving it again, or as .docx or .txt."
        ) from e
    if not text:
        raise ValueError("No text found in this file. If it's a scanned image, upload a .docx or .txt version.")
    return text


def _extract(name: str, data: bytes) -> str:
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if name.endswith(".docx"):
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return "\n".join(parts).strip()
    return data.decode("utf-8", errors="replace").strip()


STRUCTURE_PROMPT = """Convert the resume below into JSON matching this schema. Copy facts exactly; never
invent, embellish or infer anything that is not stated. Leave unknown fields empty.
For target_titles, use the job titles the person has held or clearly aims for.
Schema (JSON Schema):
{schema}

RESUME:
{resume}"""


async def structure_resume(text: str) -> CandidateProfile:
    client = ChatClient(LLMConfig.for_drafting())
    schema = json.dumps(CandidateProfile.model_json_schema())
    data = await client.chat_json(
        [{"role": "user", "content": STRUCTURE_PROMPT.format(schema=schema, resume=text[:30000])}],
        temperature=0.0,
        max_tokens=4000,
    )
    if not isinstance(data, dict):
        raise ValueError("Model returned an unexpected profile format")
    cleaned = _clean(data)
    try:
        return CandidateProfile.model_validate(cleaned)
    except ValidationError as e:
        # Local models sometimes get one field's type wrong ("6+ years"): keep everything else.
        bad = {str(err["loc"][0]) for err in e.errors() if err.get("loc")}
        return CandidateProfile.model_validate({k: v for k, v in cleaned.items() if k not in bad})


def _clean(data: dict) -> dict:
    fields = CandidateProfile.model_fields
    out = {}
    for key, value in data.items():
        if key not in fields or value is None:
            continue
        out[key] = value
    if out.get("remote_preference") not in ("remote", "hybrid", "onsite", "any"):
        out.pop("remote_preference", None)
    return out
